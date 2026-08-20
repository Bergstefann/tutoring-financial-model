"""Extract and anonymise lesson/rate data from the exported "Portfolio Lesson
Schedule" workbook.

Source: data/raw/Portfolio Lesson Schedule.xlsx (manually exported from Google
Sheets — see docs/assumptions.md for provenance). Not committed to git.

Output (both under data/, gitignored):
  - pseudonym_map.json   student first name -> "Student NN" (order of
                         appearance in the Student Config sheet)
  - lessons_anonymized.csv  student_pseudonym, lesson_date, attended,
                         duration_minutes, rate

No parent name, parent email, or student surname is read past this script —
they are dropped at the source and never reach any output file.
"""

import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "raw" / "Portfolio Lesson Schedule.xlsx"
PSEUDONYM_MAP_PATH = ROOT / "data" / "pseudonym_map.json"
LESSONS_CSV_PATH = ROOT / "data" / "lessons_anonymized.csv"

# Confirmed with the business owner (2026-08-20): every lesson slot is 30
# minutes. Not present in the source data — this is a supplied assumption,
# not an extracted fact, and must be labelled as such in the workbook.
DURATION_MINUTES = 30

# (name_col, attendance_col) pairs within each weekly block.
DAY_COLUMN_PAIRS = [(1, 2), (3, 4), (5, 6), (7, 8), (9, 10), (11, 12), (13, 14)]
BLOCK_HEIGHT = 9  # 1 header row (dates) + 8 student rows


def load_students(ws):
    """Read the Student Config sheet: ordered list of (first_name, rate)."""
    students = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        first_name, _parent_email, rate, _parent_name, _last_name = row[:5]
        if first_name is None:
            continue
        if first_name in seen:
            raise ValueError(f"Duplicate student name in Student Config: {first_name!r}")
        seen.add(first_name)
        students.append((first_name, float(rate)))
    return students


def load_lessons(ws, known_names):
    """Read the Lesson Schedule sheet: list of (first_name, date, attended)."""
    records = []
    row = 1
    while row <= ws.max_row:
        header_has_date = any(
            ws.cell(row=row, column=name_col).value is not None
            for name_col, _ in DAY_COLUMN_PAIRS
        )
        if not header_has_date:
            break

        for name_col, att_col in DAY_COLUMN_PAIRS:
            date_val = ws.cell(row=row, column=name_col).value
            if date_val is None:
                continue
            lesson_date = date_val.date()
            for student_row in range(row + 1, row + BLOCK_HEIGHT):
                name = ws.cell(row=student_row, column=name_col).value
                attended = ws.cell(row=student_row, column=att_col).value
                if name is None:
                    continue
                if name not in known_names:
                    raise ValueError(
                        f"Row {student_row}: student {name!r} not found in "
                        f"Student Config"
                    )
                if attended not in ("Y", "N"):
                    raise ValueError(
                        f"Row {student_row}: unexpected attendance value "
                        f"{attended!r} for {name!r} on {lesson_date}"
                    )
                records.append((name, lesson_date, attended))

        row += BLOCK_HEIGHT
    return records


def main():
    if not SOURCE.exists():
        raise SystemExit(f"Source file not found: {SOURCE}")

    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    students = load_students(wb["Student Config"])
    known_names = {name for name, _ in students}

    pseudonym_map = {
        name: f"Student {i + 1:02d}" for i, (name, _rate) in enumerate(students)
    }
    rate_by_name = dict(students)

    lessons = load_lessons(wb["Lesson Schedule"], known_names)

    PSEUDONYM_MAP_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(PSEUDONYM_MAP_PATH, "w", encoding="utf-8") as f:
        json.dump(pseudonym_map, f, indent=2)

    with open(LESSONS_CSV_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(
            ["student_pseudonym", "lesson_date", "attended", "duration_minutes", "rate"]
        )
        for name, lesson_date, attended in lessons:
            writer.writerow(
                [
                    pseudonym_map[name],
                    lesson_date.isoformat(),
                    attended,
                    DURATION_MINUTES,
                    rate_by_name[name],
                ]
            )

    delivered = sum(1 for _, _, a in lessons if a == "Y")
    print(f"Students: {len(students)}")
    print(f"Lesson-slots in schedule: {len(lessons)}")
    print(f"Lessons actually delivered (attended=Y): {delivered}")
    print(f"Distinct dates: {len(set(d for _, d, _ in lessons))}")
    print(f"Wrote {PSEUDONYM_MAP_PATH.relative_to(ROOT)}")
    print(f"Wrote {LESSONS_CSV_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
